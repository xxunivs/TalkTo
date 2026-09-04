# 작성본 md → 보기 쉬운 워드. 꼬리표 정리 + ### 처리 + M2 카드 표로 분해.
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT="맑은 고딕"
GRAY=RGBColor(0x70,0x70,0x70); DARK=RGBColor(0x1a,0x1a,0x1a)
BLUE=RGBColor(0x1F,0x4E,0x79); WARN=RGBColor(0xB0,0x3A,0x2B)

# 지울 꼬리표 / 남길 꼬리표
SRC=r'(?:본인|직접|전언|추정)'
FMT=r'(?:단일|복수 ?\d*|단답 ?\d*|숫자|서술[^\]]*|표[^\]]*|등급|확인)'

def clean(s):
    if s.lstrip().startswith('|'): return s              # 표는 손대지 않는다
    s = re.sub(r'`\[([^\]]*)\]`', r'[\1]', s)            # 백틱 벗기기
    s = re.sub(r'\[상태: *([^\]]*)\]', r'[\1]', s)
    def keep(m):
        v=m.group(1).strip()
        if re.match(r'^(모름|확인 필요)', v): return '['+v+']'
        if re.fullmatch(rf'{SRC}(?: *\+ *{SRC})*', v): return ''      # 출처
        if re.match(r'^(답함|건너뜀|미수집)', v): return ''            # 상태
        if re.fullmatch(FMT, v): return ''                            # 형식 코드
        return m.group(0)
    s = re.sub(r'\[([^\]]*)\]', keep, s)
    s = s.replace('`','')
    s = re.sub(r'(?<!\*)\*([^*]+)\*(?!\*)', r'\1', s)              # 홑별표 이탤릭
    s = re.sub(r'\s*[—·/]\s*$', '', s.rstrip())
    s = re.sub(r'\(\s*\)|:\s*$', '', s)
    s = re.sub(r'\s*\+\s*$', '', s)
    s = re.sub(r'^\s*[—–]\s+', '', s)
    s = re.sub(r'  +', ' ', s).rstrip()
    return s

# 문항끼리 가리키는 편집 주석 — 답이 아니라 작성 관리용이라 문서에서 뺀다
XREF=re.compile(r'(경위는|자세한 것은|참고)\s*[A-Z]?\d?[TCPM]|→\s*[CMPT]\d|'
                r'[CMPT]\d[\d\-a-b]*\s*(카드|표|참조|에서|으로|에 반영|기본값)|'
                r'기준으로 고정|자동 생성|승계|합류|미기재 방침|의도된 공백')
def is_note(s):
    s=s.strip()
    if not s or s.startswith(('**','|','#','>','-','답:','□','■')): return False
    if s.startswith('메모'): return False
    if '■' in s or '□' in s: return False
    if re.match(r'^[CMPT]\d[\d\-a-b·]*\s*[:：]', s): return False
    return bool(XREF.search(s))

def shade(p,c):
    e=OxmlElement('w:shd'); e.set(qn('w:val'),'clear'); e.set(qn('w:fill'),c)
    p._p.get_or_add_pPr().append(e)
def border(p,color='1F4E79',sz='12'):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); pPr.append(b)
    e=OxmlElement('w:bottom'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz)
    e.set(qn('w:space'),'4'); e.set(qn('w:color'),color); b.append(e)

INLINE=re.compile(r'(\*\*.+?\*\*|`[^`]+`|\[(?:모름|확인 필요)[^\]]*\])')
def rich(p,text,size=10,base=None):
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**'):
            r=p.add_run(tok[2:-2]); r.bold=True; r.font.color.rgb=base or DARK
        elif tok.startswith('`') and tok.endswith('`'):
            r=p.add_run(tok[1:-1]); r.font.color.rgb=GRAY; r.font.size=Pt(size-1.5); continue
        elif tok.startswith('[') and tok.endswith(']'):
            r=p.add_run(tok); r.font.color.rgb=WARN; r.font.size=Pt(size-1.5); continue
        else:
            r=p.add_run(tok); r.font.color.rgb=base or DARK
        r.font.size=Pt(size)
    return p

def para(doc,text="",size=10,bold=False,color=None,before=0,after=4,indent=0,rich_=True):
    p=doc.add_paragraph(); f=p.paragraph_format
    f.space_before=Pt(before); f.space_after=Pt(after); f.left_indent=Cm(indent); f.line_spacing=1.25
    if text:
        if rich_ and not bold: rich(p,text,size,color)
        else:
            r=p.add_run(text); r.font.size=Pt(size); r.bold=bold; r.font.color.rgb=color or DARK
    return p

def table(doc,rows,widths=None,head=True):
    t=doc.add_table(rows=0,cols=len(rows[0])); t.style="Table Grid"
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,v in enumerate(row):
            if widths and j<len(widths): cells[j].width=Cm(widths[j])
            p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(2)
            p.paragraph_format.line_spacing=1.15
            if head and i==0:
                r=p.add_run(v); r.bold=True; r.font.size=Pt(8.5); shade(p,'EFEFEF')
            else: rich(p,v,8.5)
    doc.add_paragraph().paragraph_format.space_after=Pt(2)

DROPPED=[]
QID=re.compile(r'^\*\*([A-Za-z0-9\-]+)\*\*\s*(.*)$')
SUB=re.compile(r'^-\s*(M\d-[0-9a-b]+(?:·[0-9a-b]+)?)\s*(?:[^:]*?)?:\s*(.*)$')

def run(src,dst,title,sub):
    doc=Document()
    for s in doc.sections:
        s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(1.8)
    st=doc.styles["Normal"]; st.font.name=FONT; st.font.size=Pt(10)
    st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)

    para(doc,title,size=20,bold=True,after=2)
    para(doc,sub,size=9,color=GRAY,after=14,bold=False,rich_=False)

    lines=open(src,encoding="utf-8").read().splitlines()
    i=next(k for k,l in enumerate(lines) if l.strip()=="---")
    lines=[clean(l) for l in lines[i+1:]]

    buf=[]; card=[]
    def flush_tbl():
        nonlocal buf
        if buf:
            rows=[[c.strip() for c in r.strip().strip('|').split('|')] for r in buf
                  if not re.match(r'^\|[\s:\|]*-[\s:\-\|]*\|$',r.strip())]
            if rows: table(doc,rows)
            buf=[]
    def flush_card():
        nonlocal card
        if card:
            table(doc,[["문항","답"]]+card,widths=[5.2,11.5])
            card=[]

    for ln in lines:
        s=ln.rstrip()
        if s.startswith('|'): flush_card(); buf.append(s); continue
        flush_tbl()
        if not s.strip() or set(s.strip())<=set('-–—_*'): continue
        m=SUB.match(s)
        if m:
            qid,val=m.group(1),m.group(2).strip()
            label=re.sub(r'^-\s*'+re.escape(qid)+r'\s*','',s).split(':')[0].strip()
            # "M2-01a: x / M2-01b: y" 처럼 한 줄에 둘이면 쪼갠다
            parts=re.split(r'\s+/\s+(?=M\d-)', qid+(" "+label if label else "")+": "+val)
            for pt in parts:
                mm=re.match(r'^(M\d-[0-9a-b·]+)\s*(.*?):\s*(.*)$', pt.strip())
                if mm: card.append([f"{mm.group(1)} {mm.group(2)}".strip(), mm.group(3).strip()])
                else:  card.append(["", pt.strip()])
            continue
        if is_note(s):
            DROPPED.append(s); continue
        flush_card()
        if s.startswith('#'):
            lvl=len(s)-len(s.lstrip('#')); txt=s.lstrip('#').strip()
            if lvl<=2:
                p=para(doc,txt,size=15,bold=True,color=BLUE,before=22,after=8); border(p)
            else:
                para(doc,txt,size=12,bold=True,color=BLUE,before=16,after=6)
            continue
        m=QID.match(s)
        if m:
            qid,txt=m.group(1),m.group(2).strip()
            p=doc.add_paragraph(); f=p.paragraph_format
            f.space_before=Pt(12); f.space_after=Pt(5); f.left_indent=Cm(0.15); f.right_indent=Cm(0.15)
            f.line_spacing=1.2; shade(p,'F2F5F8')
            r=p.add_run(qid); r.bold=True; r.font.size=Pt(10.5)
            if txt: p.add_run("  ").font.size=Pt(10.5); rich(p,txt,10.5)
            continue
        if s.startswith('답:'):
            v=s[2:].strip().lstrip(':').strip()
            v=re.sub(r'\s*[—–]\s*[^—–]*[CMPT]\d[\d\-a-b·]*\s*(카드|표|참조|에서|으로|에 반영|기본값)[^—–]*$','',v)
            para(doc,v,size=10,indent=0.55,after=4); continue
        if s.startswith('> '): para(doc,s[2:],size=8.5,color=GRAY,indent=0.55,after=2); continue
        if s.lstrip().startswith('- '): para(doc,s.lstrip()[2:],size=10,indent=0.55,after=2); continue
        if s.startswith('**') and s.endswith('**'):
            para(doc,s.strip('*'),size=10.5,bold=True,before=10,after=3); continue
        para(doc,s,size=10,indent=0.35,after=3)
    flush_tbl(); flush_card()
    doc.save(dst); print("saved:",dst)
    if DROPPED:
        print("  뺀 주석 %d줄:" % len(DROPPED))
        for x in DROPPED: print("    -", x[:100])
        DROPPED.clear()

run(*sys.argv[1:5])
