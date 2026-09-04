# 시험지 md → 워드. 가이드·예시는 질문 칸(음영 박스) 안에 넣는다.
import re, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT="맑은 고딕"
GRAY=RGBColor(0x6E,0x6E,0x6E); DARK=RGBColor(0x1a,0x1a,0x1a); BLUE=RGBColor(0x1F,0x4E,0x79)
RED=RGBColor(0xB0,0x3A,0x2B); FADE=RGBColor(0xB4,0xB4,0xB4)

def tbl_border(t, color):
    tblPr=t._tbl.tblPr; el=OxmlElement('w:tblBorders')
    for side in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+side); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4')
        e.set(qn('w:space'),'0'); e.set(qn('w:color'),color); el.append(e)
    tblPr.append(el)

def ruled(p, width_cm=None):
    """문단 아래에 줄을 긋는다 — 글자에 붙는 밑줄이 아니라 띄운 선."""
    f=p.paragraph_format
    f.left_indent=Cm(0.4)
    if width_cm: f.right_indent=Cm(max(0, 17.6-0.4-width_cm))
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); pPr.append(b)
    e=OxmlElement('w:bottom'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'6')
    e.set(qn('w:space'),'4'); e.set(qn('w:color'),'D9D9D9' if DIM[0] else '9A9A9A')
    b.append(e)

def box(doc, value=None, rows=1, blanks=None):
    """답 자리. 값이 있으면 줄 위에 얹고, 없으면 빈 줄만 긋는다."""
    fmt = CURFMT[0] or ''
    w = None if ('서술' in fmt or not fmt) else 8.0
    # 답 영역 줄 수를 시험지와 같게 (답 한 줄 + 남은 빈 줄)
    n = (1 + len(blanks or [])) if value else max(1, rows)
    for i in range(n):
        p=doc.add_paragraph(); f=p.paragraph_format
        f.space_before=Pt(0); f.space_after=Pt(LINE_AFTER+2); f.line_spacing=1.15
        ruled(p, w)
        if i==0 and value: rich(p, value, 10)
        else: p.add_run(" ").font.size=Pt(10)

def shade(p,c):
    e=OxmlElement('w:shd'); e.set(qn('w:val'),'clear'); e.set(qn('w:fill'),c)
    p._p.get_or_add_pPr().append(e)
def rule(p,color='1F4E79',sz='12'):
    b=OxmlElement('w:pBdr'); p._p.get_or_add_pPr().append(b)
    e=OxmlElement('w:bottom'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),sz)
    e.set(qn('w:space'),'4'); e.set(qn('w:color'),color); b.append(e)

DIM=[False]; CURFMT=[None]
Q_BEFORE, Q_AFTER = 15, 5      # 문항 머리
LINE_AFTER        = 4          # 선지·답·본문 한 줄
GUIDE_AFTER       = 3          # 박스 밖 안내
H1_BEFORE, H1_AFTER = 26, 9    # PART 제목
H2_BEFORE, H2_AFTER = 18, 7    # 모듈 제목
INLINE=re.compile(r'(\*\*.+?\*\*|`[^`]+`|<u>.+?</u>)')
def rich(p,text,size=10,base=None,bold=False):
    if DIM[0]: base = FADE
    for tok in INLINE.split(text):
        if not tok: continue
        if tok.startswith('**') and tok.endswith('**'):
            r=p.add_run(tok[2:-2]); r.bold=True
        elif tok.startswith('<u>') and tok.endswith('</u>'):
            r=p.add_run(tok[3:-4])
        elif tok.startswith('`') and tok.endswith('`'):
            r=p.add_run(tok[1:-1]); r.font.size=Pt(size-1.5); r.font.color.rgb=GRAY; continue
        else:
            r=p.add_run(tok); r.bold=bold
        r.font.size=Pt(size); r.font.color.rgb=base or DARK

def para(doc,text="",size=10,bold=False,color=None,before=0,after=4,indent=0):
    p=doc.add_paragraph(); f=p.paragraph_format
    f.space_before=Pt(before); f.space_after=Pt(after); f.left_indent=Cm(indent); f.line_spacing=1.3
    if text: rich(p,text,size,color,bold)
    return p

def table(doc,rows):
    n=len(rows[0])
    rows=[r[:n]+['']*(n-len(r)) for r in rows]     # 열 수를 머리행에 맞춘다
    t=doc.add_table(rows=0,cols=n); t.style="Table Grid"
    for i,row in enumerate(rows):
        cells=t.add_row().cells
        for j,v in enumerate(row):
            p=cells[j].paragraphs[0]; p.paragraph_format.space_after=Pt(3)
            if i==0:
                r=p.add_run(v); r.bold=True; r.font.size=Pt(8.5); shade(p,'EFEFEF')
            else: rich(p,v,9)
    q=doc.add_paragraph(); q.paragraph_format.space_after=Pt(0)
    q.add_run("").font.size=Pt(1)

QID=re.compile(r'^\*\*([A-Za-z0-9\-]+) · (.*?)\*\*(?:\s+`(.+?)`)?\s*$')

doc=Document()
for s in doc.sections:
    s.left_margin=s.right_margin=Cm(2.2); s.top_margin=s.bottom_margin=Cm(1.8)
st=doc.styles["Normal"]; st.font.name=FONT; st.font.size=Pt(10)
st._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
_raw=open(sys.argv[1],encoding="utf-8").read().splitlines()
_title=next((l[2:].strip() for l in _raw if l.startswith('# ') and not l.startswith('# PART')),"Persona Intake")
para(doc,_title,size=20,bold=True,after=16)

lines=open(sys.argv[1],encoding="utf-8").read().splitlines()
_s=next((i for i,l in enumerate(lines) if l.startswith('# PART')), 0)
lines=[l for l in lines[_s:] if set(l.strip())!={'-'}]

# 문항 단위로 끊어, 그 안의 > 안내줄을 질문 칸으로 끌어올린다
blocks=[]; cur=None
for ln in lines:
    if QID.match(ln) or ln.startswith('#') or (ln.startswith('##')):
        cur={'head':ln,'body':[]}; blocks.append(cur)
    elif cur is None:
        blocks.append({'head':None,'body':[ln]}); cur=blocks[-1]
    else: cur['body'].append(ln)

for b in blocks:
    h=b['head']
    DIM[0] = any(l.strip()=='<!--dim-->' for l in b['body'])
    guides=[l[2:].strip() for l in b['body'] if l.startswith('> ')]
    body=[l for l in b['body'] if not l.startswith('> ') and l.strip()!='<!--dim-->']
    if h and h.startswith('#'): CURFMT[0]=None
    if h and h.startswith('#'):
        lvl=len(h)-len(h.lstrip('#')); txt=h.lstrip('#').strip()
        if lvl==1: p=para(doc,txt,size=15,bold=True,color=BLUE,before=H1_BEFORE,after=H1_AFTER); rule(p)
        else: para(doc,txt,size=12,bold=True,color=BLUE,before=H2_BEFORE,after=H2_AFTER)
    elif h:
        qid,txt,fmt=QID.match(h).groups(); CURFMT[0]=fmt
        p=doc.add_paragraph(); f=p.paragraph_format
        f.space_before=Pt(Q_BEFORE); f.space_after=Pt(Q_AFTER); f.left_indent=Cm(0.2); f.right_indent=Cm(0.2)
        f.line_spacing=1.25; shade(p,'F1F5F9')
        r=p.add_run(qid+"  "); r.bold=True; r.font.size=Pt(10.5)
        if DIM[0]: r.font.color.rgb=FADE
        rich(p,txt,10.5,bold=True)
        if fmt:
            r=p.add_run("   "+fmt); r.font.size=Pt(8.5)
            r.font.color.rgb=FADE if DIM[0] else GRAY
        for g in guides:
            p.add_run().add_break()
            rich(p, g, 8.5, base=(FADE if DIM[0] else (RED if g.startswith('예:') else GRAY)))
        guides=[]
    buf=[]
    def flush():
        global buf
        if buf:
            rows=[[c.strip() for c in r.strip().strip('|').split('|')] for r in buf
                  if not re.match(r'^\|[\s:\|]*-[\s:\-\|]*\|$',r.strip())]
            if rows: table(doc,rows)
            buf=[]
    pend=[]
    def flush_box():
        global pend
        if pend:
            vals=[x for x in pend if not re.fullmatch(r'_{4,}', x)]
            box(doc, " ".join(vals) if vals else None, rows=len(pend),
                blanks=[x for x in pend if re.fullmatch(r'_{4,}', x)])
            pend=[]
    for s in body:
        s=s.rstrip()
        if re.fullmatch(r'_{4,}', s.strip()) or re.fullmatch(r'<u>.*</u>', s.strip()):
            flush(); pend.append(s.strip()); continue
        flush_box()
        if s.startswith('|'): buf.append(s); continue
        flush()
        if not s.strip(): continue
        if s.startswith('- '): para(doc,s[2:],size=10,indent=0.5,after=LINE_AFTER); continue
        if s.startswith('*') and s.endswith('*') and not s.startswith('**'):
            para(doc,s.strip('*'),size=9,color=GRAY,after=LINE_AFTER); continue
        if s.startswith('**') and s.endswith('**'):
            para(doc,s.strip('*'),size=10.5,bold=True,before=Q_BEFORE-3,after=Q_AFTER); continue
        para(doc,s,size=10,indent=0.4,after=LINE_AFTER)
    flush(); flush_box()
    for g in guides:
        para(doc, g, size=8.5,
             color=(FADE if DIM[0] else (RED if g.startswith('예:') else GRAY)), indent=0.6, after=GUIDE_AFTER)

doc.save(sys.argv[2]); print("saved:",sys.argv[2])
